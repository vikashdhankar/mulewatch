"""Victim-complaint matching.

An officer uploads the unit's own NCRP / 1930 complaint export (CSV or Excel). Each row's suspect identifiers
(phone, UPI ID, bank account, Telegram/WhatsApp handle, website) are normalised the same way the detection
engine normalises identifiers in ads, then stored only as keyed hashes for matching. The rest of the row is
encrypted; only the date, amount and district stay readable for statistics. Victim names and contact details
are never imported (there is no field to map them to).

A case whose identifiers hash to the same value as a complaint's is linked to that complaint: the ad is no
longer just suspicious, it is tied to money already lost."""
import csv
import io
import re
import secrets
import time
from collections import defaultdict
from datetime import datetime

from .coordination import IGNORE_DOMAINS
from .db import Complaint, ComplaintIdent, ComplaintImport
from .extract import extract_entities
from .security import decrypt, encrypt, ident_digest

FIELDS = {
    "complaint_id": "Complaint / acknowledgement number",
    "date": "Complaint date",
    "amount": "Amount lost (₹)",
    "district": "District",
    "police_station": "Police station",
    "category": "Fraud category",
    "phone": "Suspect mobile number(s)",
    "upi": "Suspect UPI ID / VPA",
    "account": "Suspect bank account number",
    "ifsc": "Suspect IFSC",
    "url": "Suspect link / handle / website",
}
IDENT_FIELDS = ("phone", "upi", "account", "ifsc", "url")
# order matters: the first rule that fits a header wins
GUESS = [
    ("ifsc", r"ifsc"),
    ("upi", r"\bupi\b|\bvpa\b|upi\s*id"),
    ("account", r"(a/?c|account|acct)\s*(no|number|num)|bank\s*a/?c|account$"),
    ("complaint_id", r"ack|acknowledg|complaint\s*(no|number|id)|ncrp|^ref"),
    ("date", r"date|reported\s*on"),
    ("amount", r"amount|loss|lost|\brs\b|₹|inr"),
    ("district", r"district|zila"),
    ("police_station", r"police\s*station|\bps\b|thana"),
    ("category", r"categor|fraud\s*type|sub\s*type|type of"),
    ("phone", r"mobile|phone|contact\s*no|\bmob\b|whats\s*app\s*(no|number)"),
    ("url", r"url|link|website|handle|telegram|social|domain|whatsapp"),
]
VICTIM = re.compile(r"victim|complainant|name|email|address|aadhaar|pan\b", re.I)

_uploads: dict[str, dict] = {}      # token -> parsed upload waiting for the officer to confirm the mapping


class ImportError_(Exception):
    pass


# ----------------------------------------------------------------------------- parsing the file

def _rows_from_csv(data: bytes) -> list[list[str]]:
    for enc in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            text = data.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    return [r for r in csv.reader(io.StringIO(text), dialect)]


def _rows_from_xlsx(data: bytes) -> list[list[str]]:
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    ws = wb.worksheets[0]
    out = []
    for row in ws.iter_rows(values_only=True):
        out.append(["" if v is None else (v.strftime("%Y-%m-%d %H:%M") if isinstance(v, datetime) else str(v))
                    for v in row])
    return out


def parse_upload(filename: str, data: bytes, user) -> dict:
    name = (filename or "").lower()
    if len(data) > 25 * 1024 * 1024:
        raise ImportError_("File too large (max 25 MB)")
    try:
        rows = _rows_from_xlsx(data) if name.endswith((".xlsx", ".xlsm")) else _rows_from_csv(data)
    except Exception as e:
        raise ImportError_(f"Couldn't read the file ({type(e).__name__}). Upload a CSV or .xlsx export.")
    # the header is the first row with at least three filled cells (exports often start with a title row)
    start = next((i for i, r in enumerate(rows[:20]) if sum(1 for c in r if str(c).strip()) >= 3), None)
    if start is None:
        raise ImportError_("No header row found")
    headers = [str(h).strip() or f"Column {i + 1}" for i, h in enumerate(rows[start])]
    body = [r + [""] * (len(headers) - len(r)) for r in rows[start + 1:] if any(str(c).strip() for c in r)]
    if not body:
        raise ImportError_("The file has a header but no rows")
    token = secrets.token_urlsafe(12)
    for k in [k for k, v in _uploads.items() if v["at"] < time.time() - 1800]:
        _uploads.pop(k, None)
    _uploads[token] = {"filename": filename, "headers": headers, "rows": body, "user": user.username, "at": time.time()}
    return {"token": token, "filename": filename, "headers": headers, "rows": len(body),
            "preview": [r[:len(headers)] for r in body[:5]], "mapping": guess_mapping(headers), "fields": FIELDS}


def guess_mapping(headers: list[str]) -> dict:
    """{header: field or ""}. Columns about the victim are left unmapped on purpose."""
    out, used = {}, set()
    for h in headers:
        low = h.lower()
        field = ""
        if not VICTIM.search(low) or re.search(r"suspect|fraudster|accused|beneficiary", low):
            for f, rx in GUESS:
                if re.search(rx, low) and (f in IDENT_FIELDS or f not in used):
                    field = f
                    break
        out[h] = field
        if field:
            used.add(field)
    return out


# ----------------------------------------------------------------------------- normalising identifiers

def norm_values(field: str, raw: str) -> list[tuple[str, str]]:
    """[(kind, value)] in the same normal form the detector uses. kind is the matching key family."""
    raw = str(raw or "").strip()
    if not raw or raw.lower() in ("na", "n/a", "nil", "-", "none", "null"):
        return []
    out = []
    if field == "phone":
        for tok in re.split(r"[,;/|\n]+", raw):
            d = re.sub(r"\D", "", tok)
            if len(d) >= 10 and d[-10] in "6789":
                out.append(("phone", "+91" + d[-10:]))
    elif field == "upi":
        out += [("upi", m.lower()) for m in re.findall(r"[A-Za-z0-9][A-Za-z0-9.\-_]{1,60}@[A-Za-z]{2,20}", raw)]
    elif field == "account":
        for tok in re.split(r"[,;/|\n]+", raw):
            d = re.sub(r"[\s\-]", "", tok)
            if d.isdigit() and 9 <= len(d) <= 18:
                out.append(("account", d))
    elif field == "ifsc":
        out += [("ifsc", m.upper()) for m in re.findall(r"[A-Za-z]{4}0[A-Za-z0-9]{6}", raw)]
    elif field == "url":
        out += handle_keys(extract_entities(raw))
        for tok in re.split(r"[,;\s|]+", raw):           # a bare @handle or channel name
            t = tok.strip().lstrip("@")
            if re.fullmatch(r"[A-Za-z][A-Za-z0-9_]{4,31}", t) and "." not in tok:
                out.append(("handle", t.lower()))
    return list(dict.fromkeys(out))


def handle_keys(ents: dict) -> list[tuple[str, str]]:
    """Identifiers from extracted entities, in matching form. Telegram links and @handles share one
    family ("handle") so t.me/x in an ad matches @x in a complaint."""
    out = []
    for p in ents.get("phone", []):
        out.append(("phone", p))
    for w in ents.get("whatsapp", []):
        if w.startswith("+91"):
            out.append(("phone", w))
    for u in ents.get("upi", []):
        out.append(("upi", u.lower()))
    for a in ents.get("bank_account", []):
        out.append(("account", a))
    for t in ents.get("telegram", []):
        name = t.split("/", 1)[1] if "/" in t else t
        if not name.startswith("+"):
            out.append(("handle", name.lower()))
    for h in ents.get("handle", []):
        out.append(("handle", h.lstrip("@").lower()))
    for d in ents.get("domains", []):
        if d not in IGNORE_DOMAINS and not d.endswith(".gov.in"):
            out.append(("domain", d))
    return list(dict.fromkeys(out))


def _date(v: str) -> datetime | None:
    v = str(v or "").strip()
    for fmt in ("%Y-%m-%d %H:%M", "%Y-%m-%d %H:%M:%S", "%Y-%m-%d", "%d/%m/%Y %H:%M", "%d/%m/%Y %H:%M:%S",
                "%d/%m/%Y", "%d-%m-%Y %H:%M", "%d-%m-%Y", "%d.%m.%Y", "%d-%b-%Y", "%d %b %Y", "%d/%m/%y"):
        try:
            return datetime.strptime(v, fmt)
        except ValueError:
            continue
    return None


def _amount(v: str) -> float:
    d = re.sub(r"[^\d.]", "", str(v or "").replace(",", ""))
    try:
        return round(float(d), 2) if d else 0.0
    except ValueError:
        return 0.0


# ----------------------------------------------------------------------------- import

def do_import(session, token: str, mapping: dict, label: str, user) -> dict:
    up = _uploads.get(token)
    if not up:
        raise ImportError_("This upload has expired. Upload the file again.")
    mapping = {h: f for h, f in mapping.items() if f in FIELDS and h in up["headers"]}
    if not any(f in IDENT_FIELDS for f in mapping.values()):
        raise ImportError_("Map at least one suspect identifier column (mobile, UPI, account, link or IFSC)")
    col = {h: i for i, h in enumerate(up["headers"])}
    imp = ComplaintImport(label=label.strip() or up["filename"], filename=up["filename"], imported_by=user.username,
                          mapping=mapping)
    session.add(imp)
    session.flush()
    added = skipped = 0
    for n, row in enumerate(up["rows"], start=1):
        rec = defaultdict(list)
        for h, f in mapping.items():
            v = str(row[col[h]]).strip()
            if v:
                rec[f].append(v)
        idents = []
        for f in IDENT_FIELDS:
            for v in rec.get(f, []):
                idents += norm_values(f, v)
        idents = list(dict.fromkeys(idents))
        ref = " ".join(rec.get("complaint_id", [])) or f"{up['filename']}#{n}"
        ref_hash = ident_digest("complaint", ref.strip().upper())
        if not idents or session.query(Complaint.id).filter_by(ref_hash=ref_hash).first():
            skipped += 1
            continue
        when = _date(rec["date"][0]) if rec.get("date") else None
        amount = sum(_amount(v) for v in rec.get("amount", []))
        district = (rec.get("district", [""])[0] or "").strip().title()[:80]
        payload = {"complaint_id": ref, "date": rec["date"][0] if rec.get("date") else "", "amount": amount,
                   "district": district, "police_station": " ".join(rec.get("police_station", [])),
                   "category": " ".join(rec.get("category", [])),
                   "identifiers": [{"kind": k, "value": v} for k, v in idents]}
        c = Complaint(import_id=imp.id, ref_hash=ref_hash, complaint_date=when, amount=amount, district=district,
                      payload=encrypt(payload))
        session.add(c)
        session.flush()
        for kind, value in idents:
            if kind != "ifsc":           # an IFSC is a branch, not a person: shown, never matched on alone
                session.add(ComplaintIdent(complaint_id=c.id, kind=kind, digest=ident_digest(kind, value)))
        added += 1
    imp.rows, imp.skipped = added, skipped
    session.flush()
    _uploads.pop(token, None)
    return {"import_id": imp.id, "added": added, "skipped": skipped}


def delete_import(session, import_id: int) -> int:
    ids = [c.id for c in session.query(Complaint.id).filter_by(import_id=import_id).all()]
    if ids:
        session.query(ComplaintIdent).filter(ComplaintIdent.complaint_id.in_(ids)).delete(synchronize_session=False)
        session.query(Complaint).filter(Complaint.id.in_(ids)).delete(synchronize_session=False)
    session.query(ComplaintImport).filter_by(id=import_id).delete()
    return len(ids)


# ----------------------------------------------------------------------------- matching

def post_keys(post) -> list[tuple[str, str]]:
    return handle_keys(post.entities or {})


def match_posts(session, posts) -> list:
    """Refresh complaint links on these posts. Returns the posts that gained new complaint links."""
    want = defaultdict(list)          # digest -> [(post, kind, value)]
    for p in posts:
        for kind, value in post_keys(p):
            want[ident_digest(kind, value)].append((p, kind, value))
    found = defaultdict(set)          # digest -> {complaint ids}
    digests = list(want)
    for i in range(0, len(digests), 500):
        for ci in session.query(ComplaintIdent).filter(ComplaintIdent.digest.in_(digests[i:i + 500])).all():
            found[ci.digest].add(ci.complaint_id)
    amounts = {}
    all_ids = {c for ids in found.values() for c in ids}
    if all_ids:
        amounts = dict(session.query(Complaint.id, Complaint.amount).filter(Complaint.id.in_(all_ids)).all())
    per_post = defaultdict(dict)      # post id -> {complaint id: kinds}
    for d, cids in found.items():
        for p, kind, value in want[d]:
            for c in cids:
                per_post[p.id].setdefault(c, set()).add(kind)
    gained = []
    for p in posts:
        links = per_post.get(p.id, {})
        old = {m["complaint"] for m in (p.complaint_matches or [])}
        p.complaint_matches = [{"complaint": c, "kinds": sorted(k)} for c, k in sorted(links.items())]
        p.complaint_count = len(links)
        p.complaint_amount = round(sum(amounts.get(c, 0) or 0 for c in links), 2)
        if set(links) - old:
            gained.append(p)
    return gained


def inr(x: float) -> str:
    """Indian digit grouping: 1234567 -> 12,34,567."""
    n = int(round(x or 0))
    s, neg = str(abs(n)), n < 0
    if len(s) > 3:
        head, tail = s[:-3], s[-3:]
        groups = []
        while len(head) > 2:
            groups.insert(0, head[-2:])
            head = head[:-2]
        s = ",".join(([head] if head else []) + groups) + "," + tail
    return ("-" if neg else "") + s


def describe(post) -> str:
    if not post.complaint_count:
        return ""
    amt = f", ₹{inr(post.complaint_amount)} reported lost" if post.complaint_amount else ""
    return (f"Identifiers match {post.complaint_count} victim complaint{'s' if post.complaint_count != 1 else ''}"
            f"{amt}")


def complaint_details(session, post) -> list[dict]:
    """Decrypted complaint rows linked to a case (Investigating Officer and above only)."""
    ids = [m["complaint"] for m in post.complaint_matches or []]
    kinds = {m["complaint"]: m["kinds"] for m in post.complaint_matches or []}
    mine = {f"{k}:{v}" for k, v in post_keys(post)}
    out = []
    for c in session.query(Complaint).filter(Complaint.id.in_(ids)).all() if ids else []:
        d = decrypt(c.payload) or {}
        out.append({"id": c.id, "complaint_id": d.get("complaint_id", "(unreadable — wrong data key?)"),
                    "date": d.get("date", ""), "amount": c.amount, "district": c.district,
                    "police_station": d.get("police_station", ""), "category": d.get("category", ""),
                    "matched_on": kinds.get(c.id, []),
                    "identifiers": [dict(i, shared=f"{i['kind']}:{i['value']}" in mine) for i in d.get("identifiers", [])]})
    out.sort(key=lambda x: x["date"] or "", reverse=True)
    return out
