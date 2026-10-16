"""Bulk link checks: officers paste or upload the links they collected (from complaints, patrolling, tips),
and each one is fetched and scored in the background like a single "Check a link".

* Links are pulled out of pasted text, a .txt/.csv file or an Excel sheet; text in the same line/row is kept as
  a note (e.g. "from complaint 3120...").
* One link at a time, with a pause between them, so no site is hammered.
* A link that is already a case is not fetched again; the item points to the existing case.
* Alerts show in the command center as usual; Telegram gets one summary per list instead of one message
  per link."""
import io
import os
import re
import threading
import time
import traceback
from datetime import datetime

from . import audit
from .config import LEVEL_RANK
from .db import BulkItem, BulkJob, Post, SessionLocal

MAX_LINKS = int(os.getenv("BULK_MAX_LINKS", "500"))
PAUSE = float(os.getenv("BULK_PAUSE_SECONDS", "1.5"))
URL_RX = re.compile(r"(?:https?://|www\.)[^\s<>\"'\]\[)(,;]+|\b(?:t\.me|telegram\.me|wa\.me|chat\.whatsapp\.com)/[^\s<>\"',;]+", re.I)

_wake = threading.Event()
_state = {"running": False, "current": None}


class BulkError(Exception):
    pass


def _clean(u: str) -> str:
    u = u.strip().rstrip(".,;:!?")
    if not re.match(r"^[a-z]+://", u, re.I):
        u = "https://" + u
    return u[:500]


def extract_links(text: str) -> list[tuple[str, str]]:
    """[(url, note)] from free text: one note per line (the line's other words)."""
    out, seen = [], set()
    for line in (text or "").splitlines():
        found = URL_RX.findall(line)
        note = URL_RX.sub(" ", line)
        note = re.sub(r"[\s,;|\t]+", " ", note).strip(" -–:")[:300]
        for f in found:
            u = _clean(f)
            if u.lower() not in seen:
                seen.add(u.lower())
                out.append((u, note))
    return out


def links_from_file(filename: str, data: bytes) -> list[tuple[str, str]]:
    name = (filename or "").lower()
    if len(data) > 10 * 1024 * 1024:
        raise BulkError("File too large (max 10 MB)")
    if name.endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        lines = []
        for ws in wb.worksheets:
            for row in ws.iter_rows(values_only=True):
                lines.append("\t".join("" if v is None else str(v) for v in row))
        text = "\n".join(lines)
    else:
        for enc in ("utf-8-sig", "cp1252", "latin-1"):
            try:
                text = data.decode(enc)
                break
            except UnicodeDecodeError:
                continue
    return extract_links(text)


def create_job(links: list[tuple[str, str]], label: str, user) -> dict:
    if not links:
        raise BulkError("No links found. Each link should start with http://, https://, www. or t.me/")
    dropped = max(0, len(links) - MAX_LINKS)
    links = links[:MAX_LINKS]
    with SessionLocal() as s:
        job = BulkJob(label=(label or "").strip()[:160] or f"{len(links)} links", created_by=user.username,
                      created_by_id=user.id, total=len(links))
        s.add(job)
        s.flush()
        for url, note in links:
            s.add(BulkItem(job_id=job.id, url=url, note=note))
        s.commit()
        jid = job.id
    audit.log(user, "bulk.created", detail={"job": jid, "links": len(links), "dropped": dropped})
    start()
    _wake.set()
    return {"id": jid, "links": len(links), "dropped": dropped}


def cancel(job_id: int, user) -> int:
    with SessionLocal() as s:
        n = s.query(BulkItem).filter_by(job_id=job_id, status="queued").update({BulkItem.status: "cancelled"})
        job = s.get(BulkJob, job_id)
        if job and job.status in ("queued", "running"):
            job.status, job.finished_at = "cancelled", datetime.utcnow()
        s.commit()
    audit.log(user, "bulk.cancelled", detail={"job": job_id, "skipped": n})
    return n


def summary(s, job: BulkJob) -> dict:
    items = s.query(BulkItem).filter_by(job_id=job.id).all()
    posts = {p.id: p for p in s.query(Post).filter(Post.id.in_([i.post_id for i in items if i.post_id])).all()}
    counts = {"queued": 0, "done": 0, "existing": 0, "error": 0, "cancelled": 0}
    levels = {"Critical": 0, "High": 0, "Medium": 0, "Low": 0}
    for i in items:
        counts[i.status] = counts.get(i.status, 0) + 1
        p = posts.get(i.post_id)
        if p:
            levels[p.risk_level] = levels.get(p.risk_level, 0) + 1
    return {"id": job.id, "label": job.label, "by": job.created_by, "created_at": job.created_at.isoformat(),
            "finished_at": job.finished_at.isoformat() if job.finished_at else None, "status": job.status,
            "total": job.total, "counts": counts, "levels": levels,
            "checked": counts["done"] + counts["existing"] + counts["error"]}


def items(s, job_id: int) -> list[dict]:
    rows = s.query(BulkItem).filter_by(job_id=job_id).order_by(BulkItem.id).all()
    posts = {p.id: p for p in s.query(Post).filter(Post.id.in_([i.post_id for i in rows if i.post_id])).all()}
    out = []
    for i in rows:
        p = posts.get(i.post_id)
        out.append({"id": i.id, "url": i.url, "note": i.note, "status": i.status, "error": i.error,
                    "post": {"id": p.id, "risk_level": p.risk_level, "risk_score": p.risk_score,
                             "category": p.category, "complaint_count": p.complaint_count or 0,
                             "excerpt": (p.text or p.ocr_text or "")[:140]} if p else None})
    return out


# ----------------------------------------------------------------------------- worker

def _check_one(item_id: int):
    from .crawler import _parse_dt
    from .fetcher import FetchError, fetch_url
    from .pipeline import ingest, recompute_all
    with SessionLocal() as s:
        it = s.get(BulkItem, item_id)
        url = it.url
        existing = s.query(Post).filter(Post.url == url).order_by(Post.id.desc()).first()
        if existing:
            it.status, it.post_id, it.checked_at = "existing", existing.id, datetime.utcnow()
            s.commit()
            return
    try:
        f = fetch_url(url)
        if not f["text"] and not f["image_paths"]:
            raise FetchError(f["note"] or "Nothing readable at that link")
    except FetchError as e:
        _finish(item_id, "error", error=str(e)[:300])
        return
    with SessionLocal() as s:
        # the final address after redirects may already be a case
        existing = s.query(Post).filter(Post.url == f["final_url"]).order_by(Post.id.desc()).first()
        if existing:
            _finish(item_id, "existing", post_id=existing.id)
            return
        imgs = f["image_paths"]
        it = s.get(BulkItem, item_id)
        p = ingest(s, text=f["text"], platform=f["platform"], author=f["author"], url=f["final_url"],
                   image_path=imgs[0] if imgs else "", extra_images=imgs[1:], posted_at=_parse_dt(f["posted_at"]),
                   source="bulk", recompute=False, html=f.get("html", ""))
        if it.note:
            p.reasons = (p.reasons or []) + [f"Officer's note with the link: {it.note[:200]}"]
        s.commit()
        recompute_all(s, alert=True, push=False)
        it = s.get(BulkItem, item_id)
        it.status, it.post_id, it.checked_at = "done", p.id, datetime.utcnow()
        s.commit()


def _finish(item_id, status, post_id=None, error=""):
    with SessionLocal() as s:
        it = s.get(BulkItem, item_id)
        it.status, it.post_id, it.error, it.checked_at = status, post_id, error, datetime.utcnow()
        s.commit()


def _notify_done(job_id: int):
    from .alerts import alert_chat_id, alert_token, send_to_chat
    from .config import PUBLIC_BASE_URL
    from .db import User
    with SessionLocal() as s:
        job = s.get(BulkJob, job_id)
        sm = summary(s, job)
        u = s.get(User, job.created_by_id) if job.created_by_id else None
        chats = {c for c in (alert_chat_id(), u.telegram_chat_id if u else "") if c}
    audit.log(None, "bulk.finished", detail={"job": job_id, **sm["counts"], **{k.lower(): v for k, v in sm["levels"].items()}})
    flagged = sm["levels"]["Critical"] + sm["levels"]["High"]
    if alert_token() and chats and flagged:
        msg = (f"📋 <b>MuleWatch bulk check finished</b> — {job.label}\n"
               f"{sm['checked']} of {sm['total']} links checked: <b>{sm['levels']['Critical']} Critical</b>, "
               f"{sm['levels']['High']} High, {sm['counts']['error']} couldn't be opened.\n\n"
               f"🔎 {PUBLIC_BASE_URL}/bulk/{job_id}")
        for c in chats:
            send_to_chat(c, msg)


def _loop():
    while _state["running"]:
        try:
            with SessionLocal() as s:
                it = (s.query(BulkItem).join(BulkJob, BulkJob.id == BulkItem.job_id)
                      .filter(BulkItem.status == "queued", BulkJob.status.in_(("queued", "running")))
                      .order_by(BulkItem.id).first())
                if it:
                    job = s.get(BulkJob, it.job_id)
                    if job.status == "queued":
                        job.status = "running"
                        s.commit()
                    item_id, job_id = it.id, it.job_id
            if not it:
                _state["current"] = None
                _wake.wait(10)
                _wake.clear()
                continue
            _state["current"] = item_id
            try:
                _check_one(item_id)
            except Exception as e:                   # one bad link must not stop the list
                traceback.print_exc()
                _finish(item_id, "error", error=f"{type(e).__name__}: {str(e)[:200]}")
            with SessionLocal() as s:
                left = s.query(BulkItem).filter_by(job_id=job_id, status="queued").count()
                job = s.get(BulkJob, job_id)
                if not left and job.status == "running":
                    job.status, job.finished_at = "done", datetime.utcnow()
                    s.commit()
                    done = True
                else:
                    done = False
            if done:
                _notify_done(job_id)
            time.sleep(PAUSE)
        except Exception:
            traceback.print_exc()
            time.sleep(5)


def start():
    if not _state["running"]:
        _state["running"] = True
        threading.Thread(target=_loop, daemon=True, name="bulk").start()


def status() -> dict:
    return {"running": _state["running"], "current_item": _state["current"]}
