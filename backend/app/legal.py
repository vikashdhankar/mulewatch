"""Draft legal documents for a case: takedown notice, data preservation request, bank request and the
Section 63 BSA certificate, in English or Hindi, as Word or PDF.

Flow: a Jinja2 template (templates/legal, or an override in data/templates/legal) is filled with the case's
details -> the officer reviews and edits the text in the command center -> the edited text is exported.
Every export carries the watermark "DRAFT — to be reviewed and signed by the authorised officer"."""
import io
import re
from datetime import datetime, timedelta
from html import escape
from pathlib import Path

from jinja2 import Environment, FileSystemLoader, StrictUndefined

from . import banks, complaints, evidence
from .config import DATA_DIR
from .pdfkit import has_devanagari, hindi_ready, para_text, safe, style

TEMPLATES = Path(__file__).parent / "templates" / "legal"
OVERRIDES = DATA_DIR / "templates" / "legal"
WATERMARK = "DRAFT — to be reviewed and signed by the authorised officer"
WATERMARK_HI = "प्रारूप — अधिकृत अधिकारी द्वारा जाँच एवं हस्ताक्षर हेतु"

KINDS = {
    "takedown": {"en": "Takedown notice to the platform (Sec. 79(3)(b) IT Act, via Sahyog)", "hi": "प्लेटफ़ॉर्म को टेकडाउन सूचना"},
    "preservation": {"en": "Data preservation request to the platform", "hi": "डेटा सुरक्षित रखने हेतु अनुरोध"},
    "bank": {"en": "Request to the bank's nodal officer", "hi": "बैंक के नोडल अधिकारी को अनुरोध"},
    "bsa63": {"en": "Section 63 BSA certificate for the evidence", "hi": "धारा 63 बी.एस.ए. प्रमाण पत्र"},
}
KIND_LABEL = {"ifsc": ("IFSC", "आई.एफ.एस.सी."), "account": ("Account number", "खाता संख्या"),
              "upi": ("UPI ID", "यूपीआई आईडी")}
ENT_LABEL = {"phone": "Mobile number", "upi": "UPI ID", "bank_account": "Bank account number", "ifsc": "IFSC",
             "telegram": "Telegram", "whatsapp": "WhatsApp", "handle": "Handle", "email": "E-mail",
             "card_number": "Card number (masked)", "crypto_wallet": "Crypto wallet", "url": "Link"}
ENT_LABEL_HI = {"phone": "मोबाइल नंबर", "upi": "यूपीआई आईडी", "bank_account": "बैंक खाता संख्या", "ifsc": "आई.एफ.एस.सी.",
                "telegram": "टेलीग्राम", "whatsapp": "व्हाट्सऐप", "handle": "हैंडल", "email": "ई-मेल",
                "card_number": "कार्ड नंबर (आंशिक)", "crypto_wallet": "क्रिप्टो वॉलेट", "url": "लिंक"}


def _env() -> Environment:
    paths = [str(OVERRIDES)] if OVERRIDES.exists() else []
    env = Environment(loader=FileSystemLoader(paths + [str(TEMPLATES)]), undefined=StrictUndefined,
                      trim_blocks=False, lstrip_blocks=False, keep_trailing_newline=True)
    return env


def _one_line(s: str, n: int = 400) -> str:
    """Values from posts go inside the line format: no newlines, no leading markers, no [[ ]]."""
    s = re.sub(r"\s+", " ", (s or "")).strip().replace("[[", "[").replace("]]", "]").replace("|", "/")
    return s[:n] + ("…" if len(s) > n else "")


def _ist(dt) -> str:
    return (dt + timedelta(hours=5, minutes=30)).strftime("%d %b %Y, %H:%M IST") if dt else "—"


def context(session, post, user, kind: str, lang: str, bank_code: str | None = None) -> dict:
    from .db import Post
    hi = lang == "hi"
    labels = ENT_LABEL_HI if hi else ENT_LABEL
    ents = post.entities or {}
    idents = [(labels[k], _one_line(", ".join(map(str, ents[k])), 300)) for k in labels if ents.get(k)]
    platform_short = (post.platform or "the platform").split(" · ")[0]
    targets = [{"url": post.url or "[[URL]]", "author": _one_line(post.author, 80),
                "captured": _ist(post.ingested_at)}]
    # linked posts on the same platform go on the same notice
    rel_ids = [r["id"] for r in post.related or []]
    for q in session.query(Post).filter(Post.id.in_(rel_ids)).all() if rel_ids else []:
        if q.platform.split(" · ")[0] == platform_short and q.url and q.risk_level in ("High", "Critical") \
                and q.url not in {t["url"] for t in targets}:
            targets.append({"url": q.url, "author": _one_line(q.author, 80), "captured": _ist(q.ingested_at)})
    files = evidence.files_for(session, post.id)
    ev = [{"name": f.path.rsplit("/", 1)[-1], "kind": f.kind, "sha256": f.sha256, "size": f.size,
           "captured": _ist(f.captured_at), "source_url": _one_line(f.source_url, 120)} for f in files]
    comp = []
    bank_values = None
    if kind == "bank":
        g = next((x for x in banks.routes(post) if (x["bank_code"] or "?") == (bank_code or "?")), None) \
            or (banks.routes(post) or [None])[0]
        bank_values = {it["value"].lower() for it in g["items"]} if g else set()
    for c in complaints.complaint_details(session, post):
        if bank_values is not None and not any(i["shared"] and i["value"].lower() in bank_values
                                               for i in c["identifiers"]):
            continue                    # a bank letter only lists complaints that involve this bank's identifiers
        comp.append(c | {"amount_fmt": complaints.inr(c["amount"]) if c["amount"] else "—",
                         "complaint_id": _one_line(c["complaint_id"], 60), "date": _one_line(c["date"], 20)})
    bank = {"name": "", "code": "", "rows": []}
    if kind == "bank":
        groups = banks.routes(post)
        g = next((x for x in groups if (x["bank_code"] or "?") == (bank_code or "?")), groups[0] if groups else None)
        if g:
            bank = {"name": g["bank"], "code": g["bank_code"], "rows": [
                it | {"kind_label": KIND_LABEL[it["kind"]][0], "kind_label_hi": KIND_LABEL[it["kind"]][1],
                      "detail": _one_line(it["detail"], 160)} for it in g["items"]]}
    officer = {"name": getattr(user, "name", "") or "[[Officer name]]",
               "designation": getattr(user, "designation", "") or "[[Designation]]",
               "police_station": getattr(user, "police_station", "") or post.police_station or "[[Police station]]"}
    return {
        "case": f"MW-{post.id:05d}", "ref": f"MW/CYB/{datetime.now().year}/{post.id:05d}",
        "today": (datetime.utcnow() + timedelta(hours=5, minutes=30)).strftime("%d.%m.%Y"),
        "fir_number": _one_line(post.fir_number, 60), "platform_short": _one_line(platform_short, 60),
        "url": post.url or "[[URL]]", "captured_at": _ist(post.ingested_at),
        "excerpt": _one_line(post.text or post.ocr_text, 350), "identifiers": idents, "targets": targets,
        "evidence": ev, "complaints": comp, "bank": bank, "officer": officer, "system": evidence.system_info(),
    }


def render(session, post, user, kind: str, lang: str, bank_code: str | None = None) -> str:
    if kind not in KINDS or lang not in ("en", "hi"):
        raise ValueError("unknown draft")
    tpl = _env().get_template(f"{kind}_{lang}.j2")
    text = tpl.render(**context(session, post, user, kind, lang, bank_code))
    return re.sub(r"\n{3,}", "\n\n", text).strip() + "\n"


# ----------------------------------------------------------------------------- the line format

def parse(text: str) -> list[tuple]:
    blocks, lines, i = [], text.splitlines(), 0
    while i < len(lines):
        ln = lines[i].rstrip()
        if not ln.strip():
            blocks.append(("blank",))
        elif ln.startswith("# "):
            blocks.append(("title", ln[2:].strip()))
        elif ln.startswith("## "):
            blocks.append(("h2", ln[3:].strip()))
        elif ln.startswith(">> "):
            blocks.append(("right", ln[3:].strip()))
        elif ln.startswith("- "):
            blocks.append(("bullet", ln[2:].strip()))
        elif ln.strip() == "~~":
            sig = []
            while i + 1 < len(lines) and lines[i + 1].strip():
                i += 1
                sig.append(lines[i].strip())
            blocks.append(("sign", sig))
        elif ln.lstrip().startswith("|"):
            rows, header = [], False
            while i < len(lines) and lines[i].lstrip().startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if all(re.fullmatch(r":?-{2,}:?", c) for c in cells if c):
                    header = len(rows) == 1
                else:
                    rows.append(cells)
                i += 1
            i -= 1
            if rows:
                blocks.append(("table", rows, header))
        else:
            blocks.append(("para", ln.strip()))
        i += 1
    return blocks


_INLINE = re.compile(r"(\*\*.+?\*\*|\[\[.+?\]\])")


def _pdf_inline(s: str) -> str:
    out = []
    for part in _INLINE.split(safe(s)):
        if part.startswith("**") and part.endswith("**"):
            out.append("<b>" + escape(part[2:-2], quote=False) + "</b>")
        elif part.startswith("[[") and part.endswith("]]"):
            out.append('<font backColor="#ffe866">[' + escape(part[2:-2], quote=False) + "]</font>")
        else:
            out.append(escape(part, quote=False))
    return "".join(out)


# ----------------------------------------------------------------------------- PDF

def to_pdf(text: str, title: str, case: str) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER, TA_RIGHT
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    hi = has_devanagari(text) and hindi_ready()
    if not hi:                          # the built-in PDF fonts have no rupee sign
        text = text.replace("₹", "Rs. ")
    st = {
        "p": style("p", hi, fontSize=10.5, leading=15, spaceAfter=2),
        "title": style("t", hi, fontSize=13, leading=19, alignment=TA_CENTER, bold=True, spaceBefore=8, spaceAfter=8),
        "h2": style("h", hi, fontSize=11, leading=16, bold=True, spaceBefore=8, spaceAfter=3,
                    textColor=colors.HexColor("#1b2a4a")),
        "right": style("r", hi, fontSize=10, leading=14, alignment=TA_RIGHT),
        "bullet": style("b", hi, fontSize=10.5, leading=15, leftIndent=12, bulletIndent=2, spaceAfter=2),
        "cell": style("c", hi, fontSize=8.5, leading=11.5),
        "cellb": style("cb", hi, fontSize=8.5, leading=11.5, bold=True),
        "mono": style("m", False, fontName="Courier", fontSize=6.8, leading=9),
        "sign": style("s", hi, fontSize=10.5, leading=15, alignment=TA_RIGHT),
    }
    P = lambda s, k: Paragraph(para_text(_pdf_inline(s), hi, st[k].fontName.endswith("Bold")), st[k])  # noqa: E731
    W = A4[0] - 36 * mm
    story = []
    for b in parse(text):
        kind = b[0]
        if kind == "blank":
            story.append(Spacer(1, 4))
        elif kind == "table":
            rows, header = b[1], b[2]
            n = max(len(r) for r in rows)
            rows = [r + [""] * (n - len(r)) for r in rows]
            weight = []
            for c in range(n):
                longest = max(len(r[c]) for r in rows)
                weight.append(36 if any(re.fullmatch(r"[0-9a-f]{64}", r[c]) for r in rows) else min(max(longest, 4), 48))
            widths = [W * w / sum(weight) for w in weight]
            data = []
            for ri, r in enumerate(rows):
                cells = []
                for c in r:
                    if re.fullmatch(r"[0-9a-f]{64}", c):
                        cells.append(Paragraph(c[:32] + "<br/>" + c[32:], st["mono"]))
                    else:
                        cells.append(P(c, "cellb" if header and ri == 0 else "cell"))
                data.append(cells)
            t = Table(data, colWidths=widths, repeatRows=1 if header else 0)
            ts = [("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#9aa4b5")), ("VALIGN", (0, 0), (-1, -1), "TOP"),
                  ("LEFTPADDING", (0, 0), (-1, -1), 4), ("RIGHTPADDING", (0, 0), (-1, -1), 4)]
            if header:
                ts.append(("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#e6ebf3")))
            t.setStyle(TableStyle(ts))
            story += [Spacer(1, 3), t, Spacer(1, 4)]
        elif kind == "sign":
            story += [Spacer(1, 18)] + [P(s, "sign") for s in b[1]]
        elif kind == "bullet":
            story.append(Paragraph(para_text(_pdf_inline(b[1]), hi), st["bullet"], bulletText="•"))
        else:
            story.append(P(b[1], {"para": "p"}.get(kind, kind)))

    mark = WATERMARK_HI if hi else WATERMARK
    generated = (datetime.utcnow() + timedelta(hours=5, minutes=30)).strftime("%d %b %Y %H:%M IST")

    def decorate(canvas, doc):
        canvas.saveState()
        canvas.setFillColor(colors.HexColor("#c0392b"))
        canvas.setFont("Helvetica-Bold", 8.5)
        canvas.drawCentredString(A4[0] / 2, A4[1] - 9 * mm, WATERMARK)
        canvas.setFillColor(colors.Color(0.75, 0.2, 0.2, alpha=0.11))
        canvas.translate(A4[0] / 2, A4[1] / 2)
        canvas.rotate(42)
        canvas.setFont("Helvetica-Bold", 96)
        canvas.drawCentredString(0, 0, "DRAFT")
        canvas.setFont("Helvetica-Bold", 15)
        canvas.drawCentredString(0, -30, "to be reviewed and signed by the authorised officer")
        canvas.restoreState()
        canvas.saveState()
        canvas.setFont("Helvetica", 7.5)
        canvas.setFillColor(colors.HexColor("#666666"))
        canvas.drawString(18 * mm, 9 * mm, f"MuleWatch {case} · {title[:70]} · generated {generated}")
        canvas.drawRightString(A4[0] - 18 * mm, 9 * mm, f"Page {doc.page}")
        canvas.restoreState()

    buf = io.BytesIO()
    SimpleDocTemplate(buf, pagesize=A4, leftMargin=18 * mm, rightMargin=18 * mm, topMargin=16 * mm,
                      bottomMargin=16 * mm, title=f"{title} — {case} (DRAFT)", author="MuleWatch",
                      subject=mark).build(story, onFirstPage=decorate, onLaterPages=decorate)
    return buf.getvalue()


# ----------------------------------------------------------------------------- Word

_VML = """<w:p xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"
 xmlns:v="urn:schemas-microsoft-com:vml" xmlns:o="urn:schemas-microsoft-com:office:office"><w:r><w:pict>
<v:shapetype id="_x0000_t136" coordsize="21600,21600" o:spt="136" adj="10800" path="m@7,l@8,m@5,21600l@6,21600e">
<v:formulas><v:f eqn="sum #0 0 10800"/><v:f eqn="prod #0 2 1"/><v:f eqn="sum 21600 0 @1"/><v:f eqn="sum 0 0 @2"/>
<v:f eqn="sum 21600 0 @3"/><v:f eqn="if @0 @3 0"/><v:f eqn="if @0 21600 @1"/><v:f eqn="if @0 0 @2"/>
<v:f eqn="if @0 @4 21600"/><v:f eqn="mid @5 @6"/><v:f eqn="mid @8 @5"/><v:f eqn="mid @7 @8"/><v:f eqn="mid @6 @7"/>
<v:f eqn="sum @6 0 @5"/></v:formulas><v:path textpathok="t" o:connecttype="custom"
 o:connectlocs="@9,0;@10,10800;@11,21600;@12,10800" o:connectangles="270,180,90,0"/>
<v:textpath on="t" fitshape="t"/><v:handles><v:h position="#0,bottomRight" xrange="6629,14971"/></v:handles>
<o:lock v:ext="edit" text="t" shapetype="t"/></v:shapetype>
<v:shape id="MuleWatchDraft" o:spid="_x0000_s2049" type="#_x0000_t136"
 style="position:absolute;margin-left:0;margin-top:0;width:460pt;height:115pt;rotation:315;z-index:-251654144;mso-position-horizontal:center;mso-position-horizontal-relative:margin;mso-position-vertical:center;mso-position-vertical-relative:margin"
 o:allowincell="f" fillcolor="#e6b8b7" stroked="f"><v:fill opacity=".45"/>
<v:textpath style="font-family:&quot;Calibri&quot;;font-size:1pt" string="DRAFT"/></v:shape></w:pict></w:r></w:p>"""


def to_docx(text: str, title: str, case: str) -> bytes:
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX
    from docx.oxml import parse_xml
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor

    hi = has_devanagari(text)
    doc = Document()
    sec = doc.sections[0]
    sec.page_height, sec.page_width = Pt(842), Pt(595)
    sec.left_margin = sec.right_margin = Pt(51)
    sec.top_margin = sec.bottom_margin = Pt(48)
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    rfonts = normal.element.get_or_add_rPr().get_or_add_rFonts()
    rfonts.set(qn("w:cs"), "Nirmala UI")         # Devanagari in Word on Windows; LibreOffice falls back to Noto
    normal.element.get_or_add_rPr().append(parse_xml(
        '<w:szCs xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" w:val="22"/>'))
    normal.paragraph_format.space_after = Pt(2)

    head = sec.header.paragraphs[0]
    head.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = head.add_run(WATERMARK)
    r.bold, r.font.size, r.font.color.rgb = True, Pt(8.5), RGBColor(0xC0, 0x39, 0x2B)
    sec.header._element.append(parse_xml(_VML))
    foot = sec.footer.paragraphs[0]
    fr = foot.add_run(f"MuleWatch {case} · {title} · {WATERMARK}")
    fr.font.size, fr.font.color.rgb = Pt(7.5), RGBColor(0x66, 0x66, 0x66)

    def runs(par, s, bold=False, size=None):
        for part in _INLINE.split(s):
            if not part:
                continue
            if part.startswith("**") and part.endswith("**"):
                rr = par.add_run(part[2:-2])
                rr.bold = True
            elif part.startswith("[[") and part.endswith("]]"):
                rr = par.add_run("[" + part[2:-2] + "]")
                rr.font.highlight_color = WD_COLOR_INDEX.YELLOW
                rr.bold = bold
            else:
                rr = par.add_run(part)
                rr.bold = bold
            if size:
                rr.font.size = Pt(size)

    for b in parse(text):
        kind = b[0]
        if kind == "blank":
            continue
        if kind == "table":
            rows, header = b[1], b[2]
            n = max(len(x) for x in rows)
            t = doc.add_table(rows=len(rows), cols=n)
            t.style = "Table Grid"
            t.alignment = WD_TABLE_ALIGNMENT.CENTER
            for ri, row in enumerate(rows):
                for ci in range(n):
                    cell = t.cell(ri, ci)
                    val = row[ci] if ci < len(row) else ""
                    mono = bool(re.fullmatch(r"[0-9a-f]{64}", val))
                    p = cell.paragraphs[0]
                    runs(p, val, bold=header and ri == 0, size=7 if mono else 9)
                    if mono:
                        for rr in p.runs:
                            rr.font.name = "Consolas"
            doc.add_paragraph()
            continue
        if kind == "sign":
            doc.add_paragraph()
            for s in b[1]:
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                runs(p, s)
            continue
        if kind == "bullet":
            p = doc.add_paragraph(style="List Bullet")
            runs(p, b[1])
            continue
        p = doc.add_paragraph()
        if kind == "title":
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before, p.paragraph_format.space_after = Pt(8), Pt(8)
            runs(p, b[1], bold=True, size=13)
        elif kind == "h2":
            p.paragraph_format.space_before = Pt(8)
            runs(p, b[1], bold=True, size=11.5)
        elif kind == "right":
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            runs(p, b[1])
        else:
            runs(p, b[1])
    doc.core_properties.title = f"{title} — {case} (DRAFT)"
    doc.core_properties.subject = WATERMARK_HI if hi else WATERMARK
    doc.core_properties.author = "MuleWatch"
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def export(text: str, fmt: str, kind: str, lang: str, case: str) -> tuple[bytes, str, str]:
    title = KINDS[kind][lang] if kind in KINDS else "Draft"
    name = f"{case}-{kind}-{lang}-DRAFT"
    if fmt == "pdf":
        return to_pdf(text, KINDS[kind]["en"], case), f"{name}.pdf", "application/pdf"
    return (to_docx(text, title, case), f"{name}.docx",
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document")
